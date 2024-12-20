import requests
from bs4 import BeautifulSoup
import pandas as pd
import time
import os
import re
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# Retry function to handle unstable connections with user prompt at the end
def fetch_with_retry(url, retries=3):
    """Attempt to fetch a URL with retries in case of failure."""
    attempt = 0
    wait_times = [10, 20, 30]  # Retry waits in seconds
    while attempt < retries:
        try:
            response = requests.get(url)
            if response.status_code == 200:
                return response
            else:
                print(f"Failed to retrieve the page. Status code: {response.status_code}")
                raise Exception("Non-200 status code")
        except (requests.exceptions.RequestException, Exception) as e:
            attempt += 1
            print(f"Attempt {attempt} failed: {e}")
            if attempt < retries:
                wait_time = wait_times[attempt - 1]  # Get the corresponding wait time for this attempt
                print(f"Retrying in {wait_time} seconds...")
                time.sleep(wait_time)
            else:
                print(f"Failed after {retries} attempts.")

                # Prompt the user to continue or stop
                while True:
                    user_input = input("Do you want to continue trying? (y/n): ").lower()
                    if user_input == 'y':
                        print("Retrying...")
                        return fetch_with_retry(url, retries)  # Retry the entire process
                    elif user_input == 'n':
                        print("Exiting program.")
                        return None  # Return None to indicate failure
                    else:
                        print("Invalid input. Please enter 'y' or 'n'.")
    return None

# Extract CID from a phytochemical's detailed page
def extract_cid(identifier):
    detailed_page_url = f"https://cb.imsc.res.in/imppat/phytochemical-detailedpage/{identifier}"
    response = fetch_with_retry(detailed_page_url)
    if response is None:
        return None

    # Parse the page content
    soup = BeautifulSoup(response.content, 'html.parser')
    external_identifiers_section = soup.find(string="External chemical identifiers:")
    if external_identifiers_section:
        cid_tag = external_identifiers_section.find_next('a')
        if cid_tag:
            cid_text = cid_tag.text.strip()
            cid_match = re.search(r'CID:(\d+)', cid_text)
            if cid_match:
                cid = cid_match.group(1)
                print(f"Extracted CID: {cid}")
                return cid

    print("CID not found")
    return None

# Retrieve the SMILES string for a compound using PubChem API
def get_smiles_from_pubchem(cid):
    pubchem_url = f"https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/cid/{cid}/property/CanonicalSMILES/TXT"
    response = fetch_with_retry(pubchem_url)
    if response is None:
        return "N/A"

    return response.text.strip()

# Download the 3D structure file for a compound in the selected format
def download_structure(identifier, structure_folder, file_format):
    # Define the URL pattern based on the chosen file format
    if file_format == "SDF":
        structure_url = f"https://cb.imsc.res.in/imppat/images/3D/SDF/{identifier}_3D.sdf"
        extension = ".sdf"
    elif file_format == "MOL":
        structure_url = f"https://cb.imsc.res.in/imppat/images/3D/MOL/{identifier}_3D.mol"
        extension = ".mol"
    elif file_format == "PDB":
        structure_url = f"https://cb.imsc.res.in/imppat/images/3D/PDB/{identifier}_3D.pdb"
        extension = ".pdb"
    elif file_format == "PDBQT":
        structure_url = f"https://cb.imsc.res.in/imppat/images/3D/PDBQT/{identifier}_3D.pdbqt"
        extension = ".pdbqt"
    else:
        print(f"Invalid file format: {file_format}")
        return "Not found"

    response = fetch_with_retry(structure_url)
    if response is None:
        return "Not found"

    # Save the file with the selected extension
    with open(os.path.join(structure_folder, f"{identifier}_3D{extension}"), 'wb') as file:
        file.write(response.content)
    print(f"Downloaded 3D structure for {identifier} in {file_format} format")
    return "Downloaded"

# Search for phytochemicals associated with a given plant
def search_plant(plant_name, structure_folder, file_format):
    plant_page_url = f"https://cb.imsc.res.in/imppat/phytochemical/{plant_name.replace(' ', '%20')}"
    response = fetch_with_retry(plant_page_url)
    if response is None:
        return []

    # Parse the page content
    soup = BeautifulSoup(response.content, 'html.parser')
    plant_info = []
    table = soup.find('table')
    if table:
        # Iterate over rows, skipping the header row
        for row in table.find_all('tr')[1:]:
            cols = row.find_all('td')
            if len(cols) > 3:
                identifier = cols[2].text.strip()
                phytochemical_name = cols[3].text.strip()
                print(f"Processing identifier: {identifier}")
                cid = extract_cid(identifier)
                if cid:
                    smiles = get_smiles_from_pubchem(cid)
                    structure_status = download_structure(identifier, structure_folder, file_format)
                    hyperlink = f"https://cb.imsc.res.in/imppat/phytochemical-detailedpage/{identifier}" if structure_status == "Not found" else ""
                    plant_info.append([identifier, phytochemical_name, cid, smiles, structure_status, hyperlink])

    return plant_info

# Save the extracted data to an Excel file
def save_to_excel(data, file_path):
    data_with_slno = [[i+1] + row for i, row in enumerate(data)]
    df = pd.DataFrame(data_with_slno, columns=['Sl No', 'Phytochemical Identifier', 'Phytochemical Name', 'CID', 'Canonical SMILES', 'Structure', 'Hyperlink'])
    df.to_excel(file_path, index=False)

    # Apply formatting and calculate summary
    wb = load_workbook(file_path)
    ws = wb.active
    fill_downloaded = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
    fill_not_found = PatternFill(start_color="FF3300", end_color="FF3300", fill_type="solid")
    downloaded_count = 0
    not_found_count = 0

    # Apply cell styles and count download statuses
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=6, max_col=6):
        for cell in row:
            if cell.value == "Downloaded":
                cell.fill = fill_downloaded
                downloaded_count += 1
            else:
                cell.fill = fill_not_found
                not_found_count += 1

    # Apply hyperlinks for missing structures
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=7, max_col=7):
        for cell in row:
            if cell.value:
                cell.hyperlink = cell.value
                cell.value = "Link"
                cell.style = "Hyperlink"

    # Add summary rows
    total_compounds = downloaded_count + not_found_count
    summary_row = ws.max_row + 2
    ws[f'A{summary_row}'] = "No of compounds downloaded:"
    ws[f'B{summary_row}'] = downloaded_count
    ws[f'A{summary_row+1}'] = "No of missing compounds:"
    ws[f'B{summary_row+1}'] = not_found_count
    ws[f'A{summary_row+2}'] = "Total no of compounds:"
    ws[f'B{summary_row+2}'] = total_compounds

    wb.save(file_path)

# Main function for processing the input and generating outputs
def main():
    # Prompt for the input Excel file path and validate it
    while True:
        excel_path = input("Enter the full path of the Excel file containing plant names: ")
        if os.path.isfile(excel_path):
            break
        else:
            print("Invalid file path. Please enter a valid path.")

    # Prompt for the output directory and validate it
    while True:
        save_location = input("Enter the location to save the files (e.g., D:/path/to/directory/): ")
        if os.path.isdir(save_location):
            break
        else:
            print("Invalid directory. Please enter a valid location.")

    # Prompt for file format choice
    print("Choose the format for 3D structure file:")
    print("1. SDF")
    print("2. MOL")
    print("3. PDB")
    print("4. PDBQT")
    while True:
        choice = input("Enter the number corresponding to your choice: ")
        if choice == "1":
            file_format = "SDF"
            break
        elif choice == "2":
            file_format = "MOL"
            break
        elif choice == "3":
            file_format = "PDB"
            break
        elif choice == "4":
            file_format = "PDBQT"
            break
        else:
            print("Invalid choice. Please enter a valid number.")

    # Load the input Excel file
    df = pd.read_excel(excel_path)
    if 'Plant Name' not in df.columns:
        print("The Excel file must contain a column named 'Plant Name'.")
        return

    # Process each plant name in the file
    for plant_name in df['Plant Name']:
        safe_plant_name = ''.join(c for c in plant_name if c.isalnum() or c in (' ', '_')).rstrip()
        plant_folder = os.path.join(save_location, safe_plant_name)
        os.makedirs(plant_folder, exist_ok=True)
        structure_folder = os.path.join(plant_folder, "3d_structure")
        os.makedirs(structure_folder, exist_ok=True)
        file_name = f"{safe_plant_name}.xlsx"
        file_path = os.path.join(plant_folder, file_name)
        
        # Get the plant data and save to Excel
        plant_info = search_plant(plant_name, structure_folder, file_format)
        if plant_info:
            save_to_excel(plant_info, file_path)
            print(f"Results saved to {file_path}")

# Run the main function
if __name__ == "__main__":
    main()
